"""Create a "Python while loops" sample lecture for testing in a fresh session.

usage: python tests/make_while_loop_data.py

Writes the same material to two places:
  data/while_loops/              ready to use: pick Lecture ID "while_loops" in the website
                                 (the persona form fills from its persona.md)
  sample_material/while_loops/   the same files, to test the upload bar with a new Lecture ID
"""

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt

from make_example_data import make_pdf, make_pptx

SLIDES = [
    ("Introduction", "Today: while loops - repeating code as long as a condition is true."),
    ("Recap: for loops", "A for loop repeats once for every item in a sequence.\n"
                         "We use it when we know how many times to repeat."),
    ("while loop syntax", "while condition:\n    body\n"
                          "Python checks the condition before every round. When it is False, the loop stops."),
    ("Counting with while", "count = 1\nwhile count <= 5:\n    print(count)\n    count = count + 1\n"
                            "The update line (count = count + 1) moves the loop towards stopping."),
    ("Infinite loops", "If the condition never becomes False, the loop never ends.\n"
                       "Forgetting the update line is the most common cause."),
    ("break", "break stops the loop immediately, even if the condition is still True.\n"
              "Useful when we find what we are looking for."),
    ("for or while?", "for: we know the number of rounds (each item in a list).\n"
                      "while: we repeat until something happens (a target score, a total)."),
]

PDF_LINES = [
    "Notes: Python while loops",
    "A while loop repeats its body as long as its condition is True.",
    "Syntax: while condition: followed by an indented body.",
    "The condition is checked before every round, including the first one.",
    "Always change something inside the loop so the condition can become False.",
    "Example: a batsman needs 50 runs. Keep adding runs until the total reaches 50.",
    "break exits the loop at once. continue skips to the next round.",
    "Common mistake: forgetting to update the counter gives an infinite loop.",
    "Use for when you know how many rounds; use while when you stop on a condition.",
]


def make_lab_docx(path):
    doc = Document()
    doc.add_heading("Lab 5: while loops", level=1)
    doc.add_heading("Aim", level=2)
    doc.add_paragraph("Practise writing while loops that stop correctly.")
    doc.add_heading("Theory", level=2)
    doc.add_paragraph("A while loop needs three things: a starting value, a condition, and an update.")

    programs = [
        ("Program 1: Countdown", "Count down from 5 to 1, then print Go!",
         ["n = 5", "while n > 0:", "    print(n)", "    n = n - 1", 'print("Go!")'],
         "Expected output: 5 4 3 2 1 Go! (one per line)"),
        ("Program 2: Reach the target", "Add runs from each over until the total reaches 50.",
         ["overs = [12, 8, 15, 6, 11, 9]", "total = 0", "i = 0", "while total < 50:",
          "    total = total + overs[i]", "    i = i + 1", 'print("Overs needed:", i, "Total:", total)'],
         "Expected output: Overs needed: 5 Total: 52"),
        ("Program 3: Stop with break", "Find the first score above 40.",
         ["scores = [12, 33, 47, 8, 51]", "i = 0", "while i < len(scores):", "    if scores[i] > 40:",
          '        print("First big score:", scores[i])', "        break", "    i = i + 1"],
         "Expected output: First big score: 47"),
    ]
    for title, aim, code, expected in programs:
        doc.add_heading(title, level=2)
        doc.add_paragraph(aim)
        for line in code:
            run = doc.add_paragraph().add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        doc.add_paragraph(expected)
    doc.save(path)


PERSONA = """# Teacher
Name: Prof. Sharma
Subject: Introduction to Python
Class: First-year B.Tech, Section A

# Absence Note
Reason: I am at a faculty workshop today.
Message to class: Please treat this session like a normal class. Try all three lab programs on your laptop tonight.

# Teaching Style
Tone: Friendly, encouraging, a little humorous
Pace: Slow and steady; repeat key ideas once
Language: Simple English, short sentences
Habits: Uses cricket examples; asks "Clear so far?" after each idea

# Lecture Structure
Order: Recap of for loops, while loop syntax, counting with while, infinite loops and break, choosing for or while
Must cover: The while condition: syntax; updating the loop variable; infinite loops and how to avoid them; break
Avoid: while-else; nested loops; recursion; reading input() from the user

# Next Class
Topic: Functions in Python
Prep: Revise while loops and try Lab 5 Program 2 without looking at the notes
"""


if __name__ == "__main__":
    targets = [Path("data/while_loops"), Path("sample_material/while_loops")]
    first = targets[0]
    first.mkdir(parents=True, exist_ok=True)
    make_pptx(first / "while_loops.pptx", SLIDES)
    make_pdf(first / "while_notes.pdf", PDF_LINES)
    make_lab_docx(first / "lab5_while.docx")
    (first / "persona.md").write_text(PERSONA, encoding="utf-8")
    for other in targets[1:]:
        other.mkdir(parents=True, exist_ok=True)
        for f in first.iterdir():
            shutil.copy(f, other / f.name)
    for t in targets:
        print(f"Wrote {t}/: " + ", ".join(sorted(p.name for p in t.iterdir())))

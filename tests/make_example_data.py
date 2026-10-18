"""Create small sample teaching files (PPTX, DOCX, PDF) in data/<lecture_id>/.

usage: python tests/make_example_data.py [lecture_id]     (default: example)
Copies templates/persona.md into the folder if it has none.
"""

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from pptx import Presentation

SLIDES = [
    ("Introduction", "Today: list comprehensions, a short way to build lists."),
    ("Recap: for loops", "A for loop repeats code for every item in a list.\nscores = []\nfor r in runs: scores.append(r * 2)"),
    ("Syntax", "[expression for item in iterable]\nExample: [r * 2 for r in runs] doubles every score."),
    ("Filtering", "[expression for item in iterable if condition]\nExample: [r for r in runs if r >= 50] keeps only fifties."),
    ("When not to use one", "If the logic needs many lines or side effects such as print, use a normal for loop."),
]

PDF_LINES = [
    "Notes: Python list comprehensions",
    "A list comprehension builds a new list from an existing iterable.",
    "Syntax: [expression for item in iterable]",
    "Filtering: add 'if condition' at the end to keep only some items.",
    "Example: squares = [n * n for n in range(5)] gives [0, 1, 4, 9, 16].",
    "A comprehension is usually faster and shorter than append in a loop.",
    "Keep comprehensions short. Readability matters more than cleverness.",
]


def make_pptx(path, slides=SLIDES):
    prs = Presentation()
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(path)


def make_docx(path):
    doc = Document()
    doc.add_heading("Lab 4: List comprehensions", level=1)
    doc.add_heading("Aim", level=2)
    doc.add_paragraph("Practise building and filtering lists with comprehensions.")
    doc.add_heading("Theory", level=2)
    doc.add_paragraph("A list comprehension has an expression, a for clause and an optional if clause.")
    doc.add_heading("Program 1", level=2)
    doc.add_paragraph("Keep only the even numbers from a list.")
    for line in ["nums = [1, 2, 3, 4, 5, 6]", "evens = [n for n in nums if n % 2 == 0]", "print(evens)"]:
        run = doc.add_paragraph().add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)
    doc.add_paragraph("Expected output: [2, 4, 6]")
    doc.save(path)


def make_pdf(path, lines=PDF_LINES):
    """A tiny hand-written one-page PDF (no extra library needed)."""
    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    text = "BT /F1 12 Tf 72 740 Td 16 TL " + " ".join(f"({esc(l)}) Tj T*" for l in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(text)} >>\nstream\n{text}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = "%PDF-1.4\n", []
    for n, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{obj}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    Path(path).write_bytes(out.encode("latin-1"))


if __name__ == "__main__":
    lecture_id = sys.argv[1] if len(sys.argv) > 1 else "example"
    folder = Path("data") / lecture_id
    folder.mkdir(parents=True, exist_ok=True)
    persona = folder / "persona.md"
    if not persona.exists():
        shutil.copy(Path("templates/persona.md"), persona)
    make_pptx(folder / "lecture.pptx")
    make_docx(folder / "lab.docx")
    make_pdf(folder / "notes.pdf")
    print(f"Sample files written to {folder}/")

"""Load the teacher's files -> chunks -> Chroma collection, plus outline.json."""

import json
import sys
from pathlib import Path

from docx import Document as DocxDocument
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_openai import OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation
from pypdf import PdfReader

import config
from persona import load_persona_or_exit

MONO_FONTS = ("courier", "consolas", "menlo", "monaco", "mono")
MAIN_TYPES = {"pptx", "pdf", "docx"}


def _first_line(text, limit=60):
    line = text.strip().splitlines()[0] if text.strip() else ""
    return line[:limit]


def load_pptx(path):
    docs = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        text = "\n".join(texts).strip()
        if text:
            meta = {"source": path.name, "type": "pptx", "slide": number}
            docs.append(Document(page_content=text, metadata=meta))
    return docs


def load_pdf(path):
    docs = []
    for number, page in enumerate(PdfReader(path).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            meta = {"source": path.name, "type": "pdf", "page": number}
            docs.append(Document(page_content=text, metadata=meta))
    return docs


def _is_code_paragraph(paragraph):
    """True if every run with text uses a monospace font."""
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        return False
    style_font = (paragraph.style.font.name or "").lower() if paragraph.style is not None else ""
    for run in runs:
        font = (run.font.name or style_font).lower()
        if not any(mono in font for mono in MONO_FONTS):
            return False
    return True


def load_docx(path):
    """One Document per heading section. Monospace paragraphs become ``` code blocks."""
    sections = []  # list of (heading, [lines])
    heading, lines, code = "Introduction", [], []

    def flush_code():
        if code:
            lines.append("```\n" + "\n".join(code) + "\n```")
            code.clear()

    for paragraph in DocxDocument(path).paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading") and paragraph.text.strip():
            flush_code()
            if lines:
                sections.append((heading, lines))
            heading, lines = paragraph.text.strip(), []
        elif _is_code_paragraph(paragraph):
            code.append(paragraph.text)
        elif paragraph.text.strip():
            flush_code()
            lines.append(paragraph.text.strip())
    flush_code()
    if lines:
        sections.append((heading, lines))

    docs = []
    for title, body in sections:
        meta = {"source": path.name, "type": "docx", "section": title}
        docs.append(Document(page_content=title + "\n" + "\n".join(body), metadata=meta))
    return docs


def load_text(path):
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    return [Document(page_content=text, metadata={"source": path.name, "type": "text"})]


LOADERS = {".pptx": load_pptx, ".pdf": load_pdf, ".docx": load_docx, ".txt": load_text, ".md": load_text}


def _outline_items(docs):
    items = []
    for d in docs:
        m = d.metadata
        if "slide" in m:
            items.append(f"{m['slide']}: {_first_line(d.page_content)}")
        elif "page" in m:
            items.append(f"page {m['page']}: {_first_line(d.page_content)}")
        elif "section" in m:
            items.append(m["section"])
    return items


def load_folder(folder):
    """Load every supported file. Returns (docs, outline). persona.md is always included."""
    docs, outline = [], []
    for path in sorted(Path(folder).iterdir()):
        if path.name.startswith(".") or path.is_dir():
            continue
        loader = LOADERS.get(path.suffix.lower())
        if loader is None:
            print(f"[index] skipping unsupported file: {path.name}")
            continue
        try:
            file_docs = loader(path)
        except Exception as e:
            print(f"[index] could not read {path.name}: {e} (skipped)")
            continue
        print(f"[index] loaded {path.name}: {len(file_docs)} part(s)")
        for d in file_docs:
            d.metadata["order"] = len(docs)
            docs.append(d)
        file_type = path.suffix.lower().lstrip(".")
        outline.append({"source": path.name, "type": file_type, "items": _outline_items(file_docs)})
    return docs, outline


def split(docs):
    splitter = RecursiveCharacterTextSplitter(chunk_size=config.CHUNK_SIZE, chunk_overlap=config.CHUNK_OVERLAP)
    chunks = splitter.split_documents(docs)
    for chunk in chunks:
        if "```" in chunk.page_content:
            chunk.metadata["has_code"] = True
        # Chroma accepts only str/int/float/bool metadata: drop anything else.
        chunk.metadata = {k: v for k, v in chunk.metadata.items() if isinstance(v, (str, int, float, bool))}
    return chunks


def get_embeddings():
    return OpenAIEmbeddings(
        model=config.EMBED_MODEL,
        base_url=config.OPENROUTER_BASE_URL,
        api_key=config.OPENROUTER_API_KEY,
        check_embedding_ctx_length=False,
    )


def get_vectorstore(collection):
    return Chroma(collection_name=collection, embedding_function=get_embeddings(), persist_directory=config.CHROMA_DIR)


def describe_location(meta):
    """'lecture.pptx slide 4' style label for a chunk."""
    label = meta.get("source", "?")
    for key in ("slide", "page", "section"):
        if key in meta:
            label += f" {key} {meta[key]}"
    return label


def run_index(lecture_id, reset=True, topic=None, duration_min=None):
    folder = config.data_dir(lecture_id)
    if not folder.is_dir():
        print(f"[index] Folder not found: {folder}. Put the lecture files there.")
        sys.exit(1)

    persona = load_persona_or_exit(folder / "persona.md")
    docs, outline = load_folder(folder)
    if not any(entry["type"] in MAIN_TYPES for entry in outline):
        print(f"[index] No PPTX, PDF or DOCX files in {folder}. Add teaching material first.")
        sys.exit(1)

    chunks = split(docs)
    if reset:
        try:
            get_vectorstore(lecture_id).delete_collection()
            print(f"[index] reset collection '{lecture_id}'")
        except Exception:
            pass  # Collection did not exist yet.

    store = get_vectorstore(lecture_id)
    # Stable ids, so a --no-reset run updates chunks instead of duplicating them.
    ids = [f"{c.metadata['source']}|{c.metadata['order']}|{i}" for i, c in enumerate(chunks)]
    total = (len(chunks) + config.EMBED_BATCH - 1) // config.EMBED_BATCH
    for n in range(total):
        start, end = n * config.EMBED_BATCH, (n + 1) * config.EMBED_BATCH
        try:
            store.add_documents(chunks[start:end], ids=ids[start:end])
        except Exception as e:
            print(f"[index] embedding failed on batch {n + 1}/{total}: {e}")
            print("[index] Check OPENROUTER_API_KEY and EMBED_MODEL in .env.")
            sys.exit(1)
        print(f"[index] embedded batch {n + 1}/{total}")

    out_dir = config.lecture_dir(lecture_id)
    (out_dir / "outline.json").write_text(json.dumps(outline, indent=2, ensure_ascii=False), encoding="utf-8")

    manifest_file = config.manifest_path(lecture_id)
    manifest = config.load_manifest(lecture_id) if manifest_file.exists() else {}
    # New material means any old script must be generated and approved again.
    manifest.update({
        "lecture_id": lecture_id,
        "topic": topic or config.TOPIC_NAME,
        "duration_min": duration_min or config.LECTURE_DURATION_MIN,
        "collection": lecture_id,
        "status": "indexed",
        "approved_at": None,
        "persona": persona,
    })
    manifest.setdefault("created_at", None)
    config.save_manifest(lecture_id, manifest)

    code_chunks = sum(1 for c in chunks if c.metadata.get("has_code"))
    print(f"[index] done: {len(outline)} files, {len(chunks)} chunks, {code_chunks} with code")
    print(f"[index] status: {manifest['status']}")


def test_query(lecture_id, query):
    results = get_vectorstore(lecture_id).similarity_search(query, k=config.TOP_K)
    if not results:
        print("[index] no results. Did you run index first?")
    for n, doc in enumerate(results, start=1):
        print(f"\n--- {n}. [{describe_location(doc.metadata)}]")
        print(doc.page_content[:500])

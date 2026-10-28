"""Load the teacher's files -> chunks -> Chroma collection, plus outline.json."""

import json
import sys
from pathlib import Path

from docx import Document as DocxDocument
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_openai import OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation
from pypdf import PdfReader

from eris import config
from eris.persona import load_persona_or_exit

MONO_FONTS = ("courier", "consolas", "menlo", "monaco", "mono")
MAIN_TYPES = {"pptx", "pdf", "docx"}


def _first_line(text, limit=60):
    line = text.strip().splitlines()[0] if text.strip() else ""
    return line[:limit]


def load_pptx(path):
    docs = []
    for number, slide in enumerate(Presentation(path).slides, start=1):
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        text = "\n".join(texts).strip()
        if text:
            meta = {"source": path.name, "type": "pptx", "slide": number}
            docs.append(Document(page_content=text, metadata=meta))
    return docs


def load_pdf(path):
    docs = []
    for number, page in enumerate(PdfReader(path).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            meta = {"source": path.name, "type": "pdf", "page": number}
            docs.append(Document(page_content=text, metadata=meta))
    return docs


def _is_code_paragraph(paragraph):
    """True if every run with text uses a monospace font."""
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        return False
    style_font = (paragraph.style.font.name or "").lower() if paragraph.style is not None else ""
    for run in runs:
        font = (run.font.name or style_font).lower()
        if not any(mono in font for mono in MONO_FONTS):
            return False
    return True


def load_docx(path):
    """One Document per heading section. Monospace paragraphs become ``` code blocks."""
    sections = []  # list of (heading, [lines])
    heading, lines, code = "Introduction", [], []

    def flush_code():
        if code:
            lines.append("```\n" + "\n".join(code) + "\n```")
            code.clear()

    for paragraph in DocxDocument(path).paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading") and paragraph.text.strip():
            flush_code()
            if lines:
                sections.append((heading, lines))
            heading, lines = paragraph.text.strip(), []
        elif _is_code_paragraph(paragraph):
            code.append(paragraph.text)
        elif paragraph.text.strip():
            flush_code()
            lines.append(paragraph.text.strip())
    flush_code()
    if lines:
        sections.append((heading, lines))

    docs = []
    for title, body in sections:
        meta = {"source": path.name, "type": "docx", "section": title}
        docs.append(Document(page_content=title + "\n" + "\n".join(body), metadata=meta))
    return docs


def load_text(path):
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    return [Document(page_content=text, metadata={"source": path.name, "type": "text"})]


LOADERS = {".pptx": load_pptx, ".pdf": load_pdf, ".docx": load_docx, ".txt": load_text, ".md": load_text}


def _outline_items(docs):
    items = []
    for d in docs:
        m = d.metadata
        if "slide" in m:
            items.append(f"{m['slide']}: {_first_line(d.page_content)}")
        elif "page" in m:
            items.append(f"page {m['page']}: {_first_line(d.page_content)}")
        elif "section" in m:
            items.append(m["section"])
    return items


def load_folder(folder):
    """Load every supported file. Returns (docs, outline). persona.md is always included.
    The deck (eris.material.deck) is marked in the outline with its slide count; a PDF deck's pages are its slides."""
    from eris.material.deck import deck_in, slide_count
    deck = deck_in(folder)
    docs, outline = [], []
    for path in sorted(Path(folder).iterdir()):
        if path.name.startswith(".") or path.is_dir():
            continue
        loader = LOADERS.get(path.suffix.lower())
        if loader is None:
            print(f"[index] skipping unsupported file: {path.name}")
            continue
        try:
            file_docs = loader(path)
        except Exception as e:
            print(f"[index] could not read {path.name}: {e} (skipped)")
            continue
        print(f"[index] loaded {path.name}: {len(file_docs)} part(s)")
        is_deck = deck is not None and path.name == deck.name
        for d in file_docs:
            d.metadata["order"] = len(docs)
            if is_deck and "page" in d.metadata:
                d.metadata["slide"] = d.metadata["page"]  # lecture notes as a PDF: each page is a slide
            docs.append(d)
        file_type = path.suffix.lower().lstrip(".")
        entry = {"source": path.name, "type": file_type, "items": _outline_items(file_docs)}
        if is_deck:
            entry.update(deck=True, count=slide_count(path))
        outline.append(entry)
    return docs, outline


def split(docs):
    splitter = RecursiveCharacterTextSplitter(chunk_size=config.CHUNK_SIZE, chunk_overlap=config.CHUNK_OVERLAP)
    chunks = splitter.split_documents(docs)
    for chunk in chunks:
        if "```" in chunk.page_content:
            chunk.metadata["has_code"] = True
        # Chroma accepts only str/int/float/bool metadata: drop anything else.
        chunk.metadata = {k: v for k, v in chunk.metadata.items() if isinstance(v, (str, int, float, bool))}
    return chunks


def get_embeddings():
    return OpenAIEmbeddings(
        model=config.EMBED_MODEL,
        base_url=config.OPENROUTER_BASE_URL,
        api_key=config.OPENROUTER_API_KEY,
        check_embedding_ctx_length=False,
    )


def get_vectorstore(collection):
    return Chroma(collection_name=collection, embedding_function=get_embeddings(), persist_directory=config.CHROMA_DIR)


def describe_location(meta):
    """'lecture.pptx slide 4' style label for a chunk."""
    label = meta.get("source", "?")
    for key in ("slide", "page", "section"):
        if key in meta:
            label += f" {key} {meta[key]}"
    return label


def run_index(lecture_id, reset=True, topic=None, duration_min=None, speaking_min=None):
    folder = config.data_dir(lecture_id)
    if not folder.is_dir():
        print(f"[index] Folder not found: {folder}. Put the lecture files there.")
        sys.exit(1)

    # The topic and length belong to this lecture and are set in Studio; a re-index keeps them.
    manifest_file = config.manifest_path(lecture_id)
    manifest = config.load_manifest(lecture_id) if manifest_file.exists() else {}
    topic = (topic or manifest.get("topic") or "").strip()
    duration_min = duration_min or manifest.get("duration_min") or config.LECTURE_DURATION_MIN
    speaking_min = speaking_min or manifest.get("speaking_min") or config.default_speaking_min(duration_min)
    persona = load_persona_or_exit(folder / "persona.md")
    docs, outline = load_folder(folder)
    if not any(entry["type"] in MAIN_TYPES for entry in outline):
        print(f"[index] No PPTX, PDF or DOCX files in {folder}. Add teaching material first.")
        sys.exit(1)

    # Pictures (photos, paintings, maps, diagrams, charts) carry no text: a vision model describes each one, and
    # the description joins the knowledge base like any other material.
    from eris.material.images import describe_images, picture_text
    pictures = describe_images(folder, lecture_id)
    for pic in pictures:
        if pic.get("description"):
            where = {"slide": pic["slide"]} if pic.get("slide") else {"section": pic["where"]}
            docs.append(Document(page_content=picture_text(pic),
                                 metadata={"source": pic["source"], "type": "picture", "order": len(docs), **where}))

    chunks = split(docs)
    if reset:
        try:
            get_vectorstore(lecture_id).delete_collection()
            print(f"[index] reset collection '{lecture_id}'")
        except Exception:
            pass  # Collection did not exist yet.

    store = get_vectorstore(lecture_id)
    # Stable ids, so a --no-reset run updates chunks instead of duplicating them.
    ids = [f"{c.metadata['source']}|{c.metadata['order']}|{i}" for i, c in enumerate(chunks)]
    total = (len(chunks) + config.EMBED_BATCH - 1) // config.EMBED_BATCH
    for n in range(total):
        start, end = n * config.EMBED_BATCH, (n + 1) * config.EMBED_BATCH
        try:
            store.add_documents(chunks[start:end], ids=ids[start:end])
        except Exception as e:
            print(f"[index] embedding failed on batch {n + 1}/{total}: {e}")
            print("[index] Check OPENROUTER_API_KEY and EMBED_MODEL in .env.")
            sys.exit(1)
        print(f"[index] embedded batch {n + 1}/{total}")

    out_dir = config.lecture_dir(lecture_id)
    # Graphs and diagrams in the material: the only graphs the script may turn into an interactive [VISUAL].
    from eris.material.figures import find_figures
    from eris.material.figures import figures_from_pictures
    figures = find_figures(folder, pictures)
    figures += figures_from_pictures(pictures)
    # How to teach each chart and diagram (a slider graph, drawn block by block, or shown on its slide): Jev
    # decides once from the figure itself, so the writer and the doubt agent both follow the same plan.
    from eris.generate.script_judge import decide_use
    decide_use(figures, topic)
    for entry in outline:
        entry["figures"] = [f"{f['kind']} {f['ref']}: {f['title']}" for f in figures if f["source"] == entry["source"]]
    (out_dir / "figures.json").write_text(json.dumps(figures, indent=2, ensure_ascii=False), encoding="utf-8")
    from collections import Counter
    uses = Counter(f.get("use") for f in figures)
    print(f"[index] how to teach them: {uses['interactive']} interactive, {uses['draw']} drawn, {uses['show']} shown on the slide")
    print(f"[index] found {sum(f['kind'] == 'graph' for f in figures)} graph(s), "
          f"{sum(f['kind'] in ('diagram', 'flow') for f in figures)} diagram(s) or step list(s) and "
          f"{sum(f['kind'] == 'picture' for f in figures)} other picture(s) in the material")
    from eris.material.figures import environments
    envs = environments(docs, figures, chunks)  # which class environments this material allows
    print(f"[index] environments: {', '.join(k for k, v in envs.items() if v)}")
    (out_dir / "outline.json").write_text(json.dumps(outline, indent=2, ensure_ascii=False), encoding="utf-8")

    # New material means any old script must be generated and approved again.
    manifest.update({
        "lecture_id": lecture_id,
        "topic": topic,
        "duration_min": duration_min,
        "speaking_min": speaking_min,
        "collection": lecture_id,
        "status": "indexed",
        "approved_at": None,
        "persona": persona,
        "environments": envs,
    })
    manifest.setdefault("created_at", None)
    config.save_manifest(lecture_id, manifest)

    code_chunks = sum(1 for c in chunks if c.metadata.get("has_code"))
    print(f"[index] done: {len(outline)} files, {len(chunks)} chunks, {code_chunks} with code")
    print(f"[index] status: {manifest['status']}")


def test_query(lecture_id, query):
    results = get_vectorstore(lecture_id).similarity_search(query, k=config.TOP_K)
    if not results:
        print("[index] no results. Did you run index first?")
    for n, doc in enumerate(results, start=1):
        print(f"\n--- {n}. [{describe_location(doc.metadata)}]")
        print(doc.page_content[:500])

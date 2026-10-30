"""Pictures in the teacher's material: pulled out of slides, PDFs and Word files, and understood by a vision model.

A painting of Napoleon, a map, a photo, a diagram or a chart saved as a picture carries no text, so reading the
files' text misses it. At index time every picture (big enough to matter) is sent once to a vision model with the
slide's text as context. It answers with what kind of picture it is, what it shows, the text written in it (OCR),
the steps of a diagram and the data of a chart. The answers are cached by the picture's hash, so re-indexing the
same material costs nothing.

The descriptions then feed everything else: the knowledge base (search finds them), the figures (a diagram picture
can be drawn as a flow, a chart picture recreated as a graph), the writer and doubt agent (what each slide shows),
and the class's Slides view, which shows the real pictures next to the slide's text.
"""

import base64
import hashlib
import io
import json
import re
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from eris import config

WEB_TYPES = {"png": "png", "jpg": "jpeg", "jpeg": "jpeg", "gif": "gif", "webp": "webp"}  # sent and shown as-is
KINDS = ("photo", "painting", "illustration", "diagram", "flowchart", "chart", "map", "screenshot", "table", "other")

VISION_PROMPT = """You are helping a teacher prepare a lecture. Look at this picture from their material.
Where it is: {where}
Text near it: {context}

Reply with JSON only (no ```), with these keys:
"kind": one of {kinds}
"description": 1-3 sentences on what the picture shows and what a student should notice (name people, places,
  events or objects when the picture or the nearby text makes them clear)
"text": every word or number written in the picture, as written ("" if none)
"steps": for a diagram or flowchart, its steps in order, e.g. ["Take the next item", "Run the body"] ([] otherwise)
"data": for a chart, its type, axes and values, e.g. "bar chart; x: A, B, C; y: hours 2, 3, 8" ("" otherwise)"""

PARTS_PROMPT = """This picture is from a teacher's material ({where}). A grid is drawn over it: the numbers along the top
are percent of the picture's width, the numbers down the left are percent of its height.
What it shows: {description}

List the parts a teacher would point at with a pen while explaining it: labelled components, regions, objects, stages
(at most {most}). Name each by the label written in the picture when it has one, else in 1-3 plain words. Give the box
that holds the part (with its label, if the label sits on it) as [left, top, right, bottom] in percent, read off the grid.
Also say whether it is decoration: a logo, a banner, a photo of a campus or building, a background or a stock image
that teaches nothing about the subject.
Reply with JSON only (no ```): {{"decoration": false, "parts": [{{"name": "...", "box": [l, t, r, b]}}]}}
A picture with nothing to point at (a portrait, a plain photo): "parts": []"""
MAX_PARTS = 8
GRID_MARGIN = 34  # pixels left of and above the picture for the grid's numbers


# ---------- pulling pictures out of the files (no LLM) ----------

def _shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == 6:  # a group: look inside
            yield from _shapes(shape.shapes)


def _pptx_images(path):
    """The pictures on each slide, except icons: a picture covering less than MIN_PICTURE of the slide is decoration
    (a bullet icon, a tick), and teaching it as a picture or chart only confuses the lesson and its checks."""
    from pptx import Presentation
    from eris.material.slide_images import MIN_PICTURE
    pres = Presentation(path)
    area = pres.slide_width * pres.slide_height
    out = []
    for number, slide in enumerate(pres.slides, start=1):
        texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        for shape in _shapes(slide.shapes):
            if getattr(shape, "shape_type", None) != 13:  # MSO_SHAPE_TYPE.PICTURE
                continue
            if shape.width and shape.height and shape.width * shape.height < MIN_PICTURE * area:
                continue  # an icon
            try:
                blob, ext = shape.image.blob, shape.image.ext.lower()
            except Exception:
                continue
            alt = (shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr") or [""])[0].strip()
            box = None  # where it sits on the slide, as fractions (for circling its parts)
            if shape.width and shape.height:
                W, H = pres.slide_width, pres.slide_height
                box = [round(v, 4) for v in (shape.left / W, shape.top / H, (shape.left + shape.width) / W,
                                             (shape.top + shape.height) / H)]
            out.append({"blob": blob, "ext": ext, "where": f"slide {number}", "ref": f"slide {number}",
                        "slide": number, "context": " / ".join(texts)[:400], "alt": alt, "box": box})
    return out


def _pdf_images(path, as_slides=False):
    """The pictures drawn on each page, except icons and flat colour (a box's shadow or fill stored as an image).
    as_slides: the PDF is the lesson's deck (eris.material.deck), so each picture sits on the slide of its page."""
    import pymupdf
    from eris.material.deck import page_pictures
    from eris.material.slide_images import MIN_PICTURE
    out = []
    with pymupdf.open(path) as pdf:
        for number, page in enumerate(pdf, start=1):
            real = page_pictures(pdf, page, MIN_PICTURE)
            text = page.get_text().strip()[:400]
            for info in page.get_image_info(xrefs=True):
                if not info.get("xref") or info["bbox"] not in real:
                    continue
                try:
                    image = pdf.extract_image(info["xref"])
                except Exception:
                    continue
                where = f"slide {number}" if as_slides else f"{path.name} page {number}"
                w, h, b = page.rect.width, page.rect.height, info["bbox"]
                out.append({"blob": image["image"], "ext": image.get("ext", "png").lower(), "where": where, "ref": where,
                            "slide": number if as_slides else None, "context": text, "alt": "",
                            "box": [round(v, 4) for v in (b[0] / w, b[1] / h, b[2] / w, b[3] / h)] if as_slides else None})
    return out


def _docx_images(path):
    from docx import Document
    doc = Document(path)
    out = []
    text = " ".join(p.text for p in doc.paragraphs[:20])[:400]
    for n, rel in enumerate([r for r in doc.part.rels.values() if "image" in r.reltype], start=1):
        try:
            part = rel.target_part
            ext = Path(part.partname).suffix.lstrip(".").lower()
            out.append({"blob": part.blob, "ext": ext, "where": f"{path.name} picture {n}",
                        "ref": f"{path.name} picture {n}", "slide": None, "context": text, "alt": ""})
        except Exception:
            continue
    return out


EXTRACTORS = {".pptx": _pptx_images, ".pdf": _pdf_images, ".docx": _docx_images}


def _web_image(blob, ext):
    """(bytes, ext) a browser and the vision model accept, or (None, None) for a tiny icon or an unreadable format.
    Other formats (bmp, tiff, ...) are converted to PNG; big pictures are scaled down."""
    from PIL import Image
    try:
        img = Image.open(io.BytesIO(blob))
        img.load()
    except Exception:
        return None, None
    if min(img.size) < config.IMAGE_MIN_SIDE:
        return None, None  # an icon, a bullet or a line: nothing to teach from
    if ext in WEB_TYPES and max(img.size) <= config.IMAGE_MAX_SIDE:
        return blob, ext
    img.thumbnail((config.IMAGE_MAX_SIDE, config.IMAGE_MAX_SIDE))
    out = io.BytesIO()
    img.convert("RGBA" if img.mode in ("RGBA", "LA", "P") else "RGB").save(out, "PNG")
    return out.getvalue(), "png"


def extract_images(folder):
    """Every picture worth looking at in the folder's files, each once (by content), with where it sits."""
    from eris.material.deck import deck_in, is_pdf
    deck = deck_in(folder)
    found, seen = [], {}  # key -> the picture as recorded (once per content)
    for path in sorted(Path(folder).iterdir()):
        extract = EXTRACTORS.get(path.suffix.lower())
        if extract is None or path.name.startswith("."):
            continue
        try:
            items = _pdf_images(path, as_slides=True) if is_pdf(deck) and path == deck else extract(path)
        except Exception as e:
            print(f"[index] could not read the pictures in {path.name}: {e}")
            continue
        for item in items:
            blob, ext = _web_image(item["blob"], item["ext"])
            if blob is None:
                continue
            key = hashlib.sha1(blob).hexdigest()[:16]
            if key in seen:
                # The same picture in a document and on a slide: keep the slide, so the class can show it.
                first = seen[key]
                if not first.get("slide") and item.get("slide"):
                    first.update({k: item.get(k) for k in ("where", "ref", "slide", "context", "alt", "box")}, source=path.name)
                continue
            seen[key] = {**item, "blob": blob, "ext": ext, "key": key, "source": path.name}
            found.append(seen[key])
    return found


# ---------- understanding them (vision model) ----------

_model = None


def _ask(prompt, blob, ext):
    """One question to the vision model about a picture; the JSON object in its reply ({} if none)."""
    global _model
    from langchain.chat_models import init_chat_model
    from langchain_core.messages import HumanMessage
    if _model is None:
        _model = init_chat_model(f"openrouter:{config.VISION_MODEL}", temperature=0, max_retries=1)
    data = base64.b64encode(blob).decode()
    message = HumanMessage(content=[
        {"type": "text", "text": prompt},
        {"type": "image_url", "image_url": {"url": f"data:image/{WEB_TYPES.get(ext, ext)};base64,{data}"}},
    ])
    reply = _model.invoke([message]).content
    if isinstance(reply, list):
        reply = "".join(p.get("text", "") if isinstance(p, dict) else str(p) for p in reply)
    match = re.search(r"\{.*\}", reply or "", re.S)
    return json.loads(match.group(0)) if match else {}


def _ask_vision(blob, ext, where, context):
    info = _ask(VISION_PROMPT.format(where=where, context=context or "(none)", kinds=", ".join(KINDS)), blob, ext)
    kind = str(info.get("kind", "other")).lower()
    return {"kind": kind if kind in KINDS else "other", "description": str(info.get("description", "")).strip(),
            "text": str(info.get("text", "")).strip(), "steps": [str(s) for s in info.get("steps") or []],
            "data": str(info.get("data", "")).strip()}


def _with_grid(blob):
    """The picture with a percent grid over it, numbered along the top and left, as PNG bytes: a vision model reads
    positions off a grid far better than it guesses coordinates."""
    from PIL import Image, ImageDraw
    img = Image.open(io.BytesIO(blob)).convert("RGB")
    img.thumbnail((1200, 1200))
    w, h = img.size
    canvas = Image.new("RGBA", (w + GRID_MARGIN, h + GRID_MARGIN), "white")
    canvas.paste(img, (GRID_MARGIN, GRID_MARGIN))
    over = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(over)
    for p in range(0, 101, 10):
        x, y = GRID_MARGIN + round((w - 1) * p / 100), GRID_MARGIN + round((h - 1) * p / 100)
        draw.line([(x, GRID_MARGIN), (x, GRID_MARGIN + h)], fill=(255, 0, 140, 90), width=1)
        draw.line([(GRID_MARGIN, y), (GRID_MARGIN + w, y)], fill=(255, 0, 140, 90), width=1)
        draw.text((x - 6, 8), str(p), fill=(200, 0, 110, 255))
        draw.text((4, y - 5), str(p), fill=(200, 0, 110, 255))
    out = io.BytesIO()
    Image.alpha_composite(canvas, over).convert("RGB").save(out, "PNG")
    return out.getvalue()


def _ask_parts(blob, where, description):
    """(decoration, parts): whether the picture only decorates the slide, and the parts a teacher points at,
    [{"name", "box": [x0, y0, x1, y1] as fractions of the picture}]."""
    info = _ask(PARTS_PROMPT.format(where=where, description=description or "(not known)", most=MAX_PARTS),
                _with_grid(blob), "png")
    parts, seen = [], set()
    for p in info.get("parts") or []:
        try:
            name = " ".join(str(p["name"]).split())[:60]
            x0, y0, x1, y1 = (min(100.0, max(0.0, float(v))) / 100 for v in p["box"])
        except (KeyError, TypeError, ValueError):
            continue
        if not name or name.lower() in seen or x1 - x0 < 0.01 or y1 - y0 < 0.01:
            continue
        seen.add(name.lower())
        parts.append({"name": name, "box": [round(x0, 4), round(y0, 4), round(x1, 4), round(y1, 4)]})
    return bool(info.get("decoration")), parts[:MAX_PARTS]


def _read_parts(path):
    """(decoration, parts) from a parts cache (a plain list before decoration was asked)."""
    data = json.loads(path.read_text(encoding="utf-8"))
    return (False, data) if isinstance(data, list) else (bool(data.get("decoration")), data.get("parts") or [])


def teaching(pictures):
    """The pictures worth teaching from: not a logo, campus photo or background (decoration)."""
    return [p for p in pictures if not p.get("decoration")]


def find_parts(pictures, out_dir, again=False):
    """Give every slide picture its "parts" (asked once, cached in <key>.parts.json): what the teacher's pen can
    circle inside it, and "decoration" when it teaches nothing. again: ask even for pictures already cached (or
    cached before decoration was asked). Never raises."""
    todo = []
    for pic in pictures:
        cached = out_dir / f"{pic['key']}.parts.json"
        if cached.exists() and not again:
            pic["decoration"], pic["parts"] = _read_parts(cached)
        elif pic.get("slide") and config.VISION_ENABLED:
            todo.append(pic)

    def ask(pic):
        try:
            blob = pic.get("blob") or (out_dir / f"{pic['key']}.{pic['ext']}").read_bytes()
            return pic, _ask_parts(blob, pic["where"], pic.get("description")), None
        except Exception as e:
            return pic, None, e

    if todo:
        print(f"[index] finding the parts of {len(todo)} slide picture(s) with {config.VISION_MODEL} ...")
        with ThreadPoolExecutor(max_workers=min(4, len(todo))) as pool:
            for pic, found, error in pool.map(ask, todo):
                if found is None:
                    print(f"[index] could not find the parts of the picture on {pic['where']}: {str(error)[:150]}")
                    continue
                pic["decoration"], pic["parts"] = found
                (out_dir / f"{pic['key']}.parts.json").write_text(
                    json.dumps({"decoration": found[0], "parts": found[1]}, indent=2), encoding="utf-8")


def add_parts(lecture_id, again=False):
    """Find the parts of an indexed lecture's slide pictures without indexing it again (again: ask afresh). Returns
    how many have parts."""
    out_dir = config.lecture_dir(lecture_id) / "images"
    pictures = load(lecture_id)
    find_parts(pictures, out_dir, again)
    (config.lecture_dir(lecture_id) / "images.json").write_text(json.dumps(pictures, indent=2), encoding="utf-8")
    return sum(1 for p in pictures if p.get("parts"))


def describe_images(folder, lecture_id):
    """Extract the folder's pictures, save them, and describe each with the vision model (cached by content).
    Returns the pictures as dicts (without the bytes) and writes images.json. Never raises."""
    out_dir = config.lecture_dir(lecture_id) / "images"
    out_dir.mkdir(parents=True, exist_ok=True)
    pictures = extract_images(folder)
    todo = []
    for pic in pictures:
        (out_dir / f"{pic['key']}.{pic['ext']}").write_bytes(pic["blob"])
        cached = out_dir / f"{pic['key']}.json"
        if cached.exists():
            pic.update(json.loads(cached.read_text(encoding="utf-8")))
        elif config.VISION_ENABLED:
            todo.append(pic)

    def describe(pic):
        try:
            return pic, _ask_vision(pic["blob"], pic["ext"], pic["where"], pic["context"]), None
        except Exception as e:
            return pic, None, e

    if todo:
        print(f"[index] describing {len(todo)} picture(s) with {config.VISION_MODEL} ...")
        with ThreadPoolExecutor(max_workers=min(4, len(todo))) as pool:
            for pic, info, error in pool.map(describe, todo):
                if info is None:
                    print(f"[index] could not describe the picture on {pic['where']}: {str(error)[:150]}")
                    continue
                pic.update(info)
                (out_dir / f"{pic['key']}.json").write_text(json.dumps(info, indent=2), encoding="utf-8")
    find_parts(pictures, out_dir)
    result = [{k: v for k, v in p.items() if k != "blob"} for p in pictures]
    (config.lecture_dir(lecture_id) / "images.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    described = sum(1 for p in result if p.get("description"))
    print(f"[index] found {len(result)} picture(s), {described} described")
    return result


def load(lecture_id):
    path = config.lecture_dir(lecture_id) / "images.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return []


def picture_text(pic):
    """A picture as searchable text for the knowledge base."""
    parts = [f"Picture on {pic['where']} ({pic.get('kind', 'picture')}): {pic.get('description', '')}"]
    if pic.get("text"):
        parts.append(f"Text in the picture: {pic['text']}")
    if pic.get("steps"):
        parts.append("Steps: " + " -> ".join(pic["steps"]))
    if pic.get("data"):
        parts.append(f"Chart data: {pic['data']}")
    return "\n".join(parts)


def describe_list(pictures):
    """The pictures as lines for prompts and the board."""
    rows = [p for p in teaching(pictures) if p.get("description")]
    if not rows:
        return "none"
    return "\n".join(f"- {p['where']} ({p.get('kind', 'picture')}): {p['description']}"
                     + (f" Text in it: {p['text'][:120]}" if p.get("text") else "")
                     + (f" Parts the pen can circle: {', '.join(part_names(p))}" if p.get("parts") else "") for p in rows)


def part_names(pic):
    return [p["name"] for p in pic.get("parts") or []]


def with_labels(pictures, lecture_id, render=True):
    """The pictures with the labels printed beside them on their slide added to their parts (by name only), so the
    writer and the doubt agent can circle a label like "Shear plane" that is slide text, not part of the picture."""
    from eris.material.slide_images import slide_labels
    try:
        labels = slide_labels(lecture_id, render)
    except Exception:
        return teaching(pictures)
    out, done = [], set()
    for p in teaching(pictures):
        n = p.get("slide")
        if n in labels and n not in done:
            done.add(n)
            have = {q["name"].lower() for q in p.get("parts") or []}
            p = {**p, "parts": (p.get("parts") or []) + [{"name": t} for t in labels[n] if t.lower() not in have]}
        out.append(p)
    return out


def slide_parts(pictures, slide_boxes=None):
    """{slide number: [{"name", "box"}]}: the parts of each slide's pictures, placed on the slide (fractions of it).
    A picture indexed before pictures kept their place takes the slide's picture box in the same order."""
    out = {}
    by_slide = {}
    for p in pictures:
        if p.get("slide") and p.get("parts"):
            by_slide.setdefault(p["slide"], []).append(p)
    for n, pics in by_slide.items():
        boxes = (slide_boxes or {}).get(str(n)) or []
        for k, p in enumerate(pics):
            box = p.get("box") or (boxes[k] if len(boxes) == len(pics) else None)
            if not box:
                continue
            bx0, by0, bx1, by1 = box
            w, h = bx1 - bx0, by1 - by0
            out.setdefault(str(n), []).extend(
                {"name": part["name"], "box": [round(bx0 + w * x0, 4), round(by0 + h * y0, 4),
                                               round(bx0 + w * x1, 4), round(by0 + h * y1, 4)]}
                for part in p["parts"] for x0, y0, x1, y1 in [part["box"]])
    return out


if __name__ == "__main__":
    import sys
    again = "--again" in sys.argv  # ask afresh, e.g. for pictures whose parts were found before decoration was asked
    for lecture in [a for a in sys.argv[1:] if a != "--again"]:  # python -m eris.material.images <lecture id>
        print(f"{lecture}: {add_parts(lecture, again)} picture(s) with parts")

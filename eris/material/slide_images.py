"""The teacher's slides as pictures, so the class sees the real deck (shapes, code boxes, charts) instead of its text.

LibreOffice turns the deck into a PDF and PyMuPDF draws each page. Beside each picture we keep where things are, as
fractions of the slide: every line of text (to highlight the line the teacher reads) and every picture or chart (to
circle it). Without LibreOffice the player keeps its text view of the slides.
"""

import json
import shutil
import subprocess
import tempfile
import threading
from pathlib import Path

from eris import config

WIDTH = 1600  # pixels across: sharp on a large screen, small enough to load quickly
MIN_PICTURE = 0.03  # share of the slide a picture or chart must cover to be circled: smaller ones are icons
_lock = threading.Lock()


def soffice():
    """The LibreOffice program, or None when it isn't installed."""
    return (shutil.which("soffice") or shutil.which("libreoffice")
            or next((p for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",) if Path(p).exists()), None))


def _deck(lecture_id):
    """The deck the slides come from: the first PPTX in the data folder, as actions._load_slides reads it."""
    return next(iter(sorted(Path(config.data_dir(lecture_id)).glob("*.pptx"))), None)


def _folder(lecture_id):
    return config.lecture_dir(lecture_id) / "slides"


def _boxes(deck):
    """{slide number: [box]}: where each picture and chart sits, as [x0, y0, x1, y1] fractions of the slide."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    pres = Presentation(deck)
    W, H = pres.slide_width, pres.slide_height

    def walk(shapes, out):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(s.shapes, out)
            elif ((s.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(s, "has_chart", False))
                  and s.width * s.height >= MIN_PICTURE * W * H):
                out.append([round(v, 4) for v in (s.left / W, s.top / H, (s.left + s.width) / W, (s.top + s.height) / H)])
        return out
    return {n: walk(slide.shapes, []) for n, slide in enumerate(pres.slides, 1)}


def _render(lecture_id, deck):
    import pymupdf
    folder = _folder(lecture_id)
    with tempfile.TemporaryDirectory() as tmp:
        # A private profile: a LibreOffice the teacher has open doesn't block the conversion.
        subprocess.run([soffice(), f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--convert-to", "pdf",
                        "--outdir", tmp, str(deck)], check=True, capture_output=True, timeout=180)
        pdf = pymupdf.open(Path(tmp) / f"{deck.stem}.pdf")
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir(parents=True)
        pictures, slides = _boxes(deck), {}
        for n, page in enumerate(pdf, 1):
            w, h = page.rect.width, page.rect.height
            page.get_pixmap(matrix=pymupdf.Matrix(WIDTH / w, WIDTH / w)).save(folder / f"{n}.png")
            lines = [{"text": "".join(s["text"] for s in line["spans"]).strip(),
                      "box": [round(line["bbox"][0] / w, 4), round(line["bbox"][1] / h, 4),
                              round(line["bbox"][2] / w, 4), round(line["bbox"][3] / h, 4)]}
                     for block in page.get_text("dict")["blocks"] if block.get("type") == 0 for line in block["lines"]]
            slides[str(n)] = {"ratio": round(w / h, 4), "lines": [l for l in lines if l["text"]],
                              "pictures": pictures.get(n, [])}
    info = {"deck": deck.name, "mtime": deck.stat().st_mtime, "slides": slides}
    (folder / "slides.json").write_text(json.dumps(info), encoding="utf-8")
    return info


def slide_images(lecture_id, render=True):
    """{slide number: {url, ratio, lines, pictures}} for the player, or {} (no deck, no LibreOffice, or it failed).
    Drawn once per deck and kept until the deck changes; render=False only reads what is already drawn."""
    deck = _deck(lecture_id)
    if not deck:
        return {}
    meta = _folder(lecture_id) / "slides.json"
    with _lock:  # one drawing at a time: the class page and an approval can ask together
        info = json.loads(meta.read_text(encoding="utf-8")) if meta.exists() else None
        if not info or info.get("deck") != deck.name or info.get("mtime") != deck.stat().st_mtime:
            if not render or not soffice():
                return {}
            try:
                info = _render(lecture_id, deck)
            except Exception as e:
                print(f"[slides] could not draw {deck.name}: {e}")
                return {}
    return {n: {**s, "url": f"/api/lectures/{lecture_id}/slides/{n}.png"} for n, s in info["slides"].items()}


def render_later(lecture_id):
    """Draw the slides in the background (after an approval), so the first class doesn't wait for them."""
    threading.Thread(target=slide_images, args=(lecture_id,), daemon=True).start()

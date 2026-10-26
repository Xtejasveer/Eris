"""How deep to teach each part of the material: the lesson's length comes from the material, not from a time the
teacher sets.

1. The writer's model splits the material into sub-topics, each with the slides it spans (one short call).
2. Code counts how much room the teacher gave each sub-topic: the slides it spans, and what is on them (charts,
   tables, code, pictures; told to Jev with the slide text).
3. Jev answers six yes/no questions about how the material explains each sub-topic. They ask only what the material
   itself does, never what students know, so they mean the same in every subject.
4. Code scores it: Jev's confidence in each yes, added up (three answers it is 70% sure of count 2.1: a hard
   cutoff would throw that away and flip parts between runs), + 1 per slide beyond the first. Under BRIEF_BELOW is
   brief, from IN_DEPTH_AT in depth, normal between. A "Must cover" item is never brief and the lesson's main idea
   is always in depth. Each depth has a fixed speaking
   budget, so an easy part can't eat the time and a hard one can't be rushed.
"""

import json
import re
from concurrent.futures import ThreadPoolExecutor

import httpx
from langchain_core.messages import HumanMessage

from eris import config

DECISIONS_URL = "https://openrouter.ai/api/alpha/decisions"
YES_AT = 0.5  # from this P(yes) an answer is named in Studio's "why" (the score uses every answer's confidence)
BRIEF_BELOW, IN_DEPTH_AT = 2.0, 3.5  # a part that is only stated scores ~1.5 (it has a case or two, little else)
DEPTHS = ("brief", "normal", "in depth")
# Speaking minutes per sub-topic of each depth; writing, code and pauses come on top in class.
DEPTH_MIN = {"brief": 0.6, "normal": 1.2, "in depth": 2.2}
FRAMING_MIN = {"Opening": 0.6, "Recap": 0.5, "Closing": 0.3}
MAX_SPEAKING_MIN = 17  # about 20 minutes of class (class time runs ~1.15x the speaking): past this, every part is
                       # shortened by the same share and Studio suggests splitting the material into two lessons
MATERIAL_CHARS = 2500  # of a sub-topic's material that Jev reads
# A section shorter than this share of its planned words goes back to the writer: stricter for in-depth parts, since
# rushing a hard part is the failure that matters (writers run ~25% short everywhere; normal and brief may).
MIN_SHARE = {"brief": 0.5, "normal": 0.5, "in depth": 0.7}
MONO_RE = re.compile(r"courier|consola|mono|menlo|code", re.I)  # a code font: the text is a program
FRAMING_RE = re.compile(r"\s*(opening|introduction|recap|summary|review|closing|conclusion|wrap)", re.I)


def _q(instructions, true, false):
    return {"type": "noul", "instructions": instructions, "criteria": {"true": true, "false": false}}


# The same six for every subject: each asks what the material does with the sub-topic, so it can be checked in its text.
QUESTIONS = {
    "example": _q(
        "Does material apply sub_topic to a specific case, such as a worked problem, a real event, a passage or a data point?",
        "The material applies it to at least one specific case.", "The material only states or describes it."),
    "sequence": _q(
        "Does material describe sub_topic as a series of steps or stages, such as a procedure, a process or a chain of events?",
        "The material lays it out as steps or stages in order.", "The material does not lay it out as steps or stages."),
    "builds_on": _q(
        "Does material explain sub_topic using another idea that is also explained in other_sub_topics?",
        "Its explanation relies on another idea from other_sub_topics.", "It is explained on its own."),
    "comparison": _q(
        "Does material compare or contrast sub_topic with something else?",
        "The material compares or contrasts it with something.", "The material makes no comparison."),
    "condition": _q(
        "Does material state a special case, exception or condition in which sub_topic works differently or does not apply?",
        "The material states such a special case, exception or condition.",
        "The material states none (a general remark or caveat is not one)."),
    "task": _q(
        "Does material ask the student to do something with sub_topic, such as calculate, write, analyse or apply it?",
        "The material asks the student to do something with it.", "The material asks the student to do nothing with it."),
}
LABELS = {"example": "a specific case", "sequence": "steps or stages", "builds_on": "builds on another part",
          "comparison": "a comparison", "condition": "an exception or condition", "task": "a task for students"}

SPLIT_PROMPT = """You plan a lesson on "{topic}" from a teacher's material. Split the material into the sub-topics a
teacher would explain one after another, in the order the material teaches them.

Rules:
- Every idea the material teaches belongs to exactly one sub-topic. Leave out title, agenda, "you're ready" and
  closing slides that teach nothing new.
- Put slides that explain the same idea together in one sub-topic; a slide that explains two ideas may be listed for
  both.
- Usually 3 to 8 sub-topics.
- main: true for the one sub-topic this lesson exists to teach (its central idea); false for the others.
- must_cover: true if the sub-topic is one of the teacher's "Must cover" items: {must_cover}
- excerpt: the material's own words for this sub-topic, copied (up to about 80 words). Needed when the material has
  no slides.

Return ONLY JSON: [{{"title": "...", "slides": [2, 3], "main": false, "must_cover": true, "excerpt": "..."}}]

# The material
{material}"""


def _slide_things(lecture_id):
    """{slide number: ["a chart", "a table", ...]}: what is on each slide besides text, from the deck."""
    from pathlib import Path
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    deck = next(iter(sorted(Path(config.data_dir(lecture_id)).glob("*.pptx"))), None)
    if not deck:
        return {}
    pres = Presentation(deck)
    area = pres.slide_width * pres.slide_height

    def walk(shapes, found):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(s.shapes, found)
            elif getattr(s, "has_chart", False):
                found.add("a chart")
            elif getattr(s, "has_table", False):
                found.add("a table")
            elif s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width * s.height >= 0.03 * area:
                found.add("a picture")  # smaller ones are icons
            elif s.has_text_frame and any(MONO_RE.search(r.font.name or "") for p in s.text_frame.paragraphs
                                          for r in p.runs if r.text.strip()):
                found.add("code")
        return found
    return {n: sorted(walk(slide.shapes, set())) for n, slide in enumerate(pres.slides, 1)}


def _json_list(text):
    m = re.search(r"\[.*\]", text, re.S)
    data = json.loads(m.group(0)) if m else []
    return [d for d in data if isinstance(d, dict) and str(d.get("title", "")).strip()]


def split_material(topic, slides, material, persona):
    """[{title, slides, main, must_cover, excerpt}] from the writer's model; [] if it fails."""
    from langchain.chat_models import init_chat_model
    if slides:
        text = "\n\n".join(f"## Slide {n}\n{t}" for n, t in sorted(slides.items(), key=lambda kv: int(kv[0])))
    else:
        text = material
    must = (persona or {}).get("Lecture Structure", {}).get("Must cover", "") or "(none given)"
    try:
        reply = init_chat_model(f"openrouter:{config.LLM_MODEL}").invoke(
            [HumanMessage(SPLIT_PROMPT.format(topic=topic, must_cover=must, material=text[:20000]))])
        content = reply.content if isinstance(reply.content, str) else "".join(
            b.get("text", "") if isinstance(b, dict) else str(b) for b in reply.content)
        parts = _json_list(content)
    except Exception as e:
        print(f"[depth] could not split the material into sub-topics: {str(e)[:200]}")
        return []
    count = len(slides)
    for p in parts:
        p["title"] = str(p["title"]).strip()
        p["slides"] = sorted({int(n) for n in p.get("slides") or [] if str(n).isdigit() and 1 <= int(n) <= count})
        p["main"], p["must_cover"] = bool(p.get("main")), bool(p.get("must_cover"))
    return parts


def _ask_jev(state):
    """{question: P(yes)} for one sub-topic, or None if Jev can't be reached."""
    try:
        resp = httpx.post(DECISIONS_URL, timeout=30, headers={"Authorization": f"Bearer {config.OPENROUTER_API_KEY}"},
                          json={"model": config.JUDGE_MODEL, "state": state, "questions": QUESTIONS})
        resp.raise_for_status()
        answers = resp.json()["answers"]
        return {k: float(answers.get(k, {}).get("noul", 0)) for k in QUESTIONS}
    except Exception as e:
        print(f"[depth] Jev could not be reached for \"{state['sub_topic']}\": {str(e)[:160]}")
        return None


def _depth(score, part):
    depth = DEPTHS[0] if score < BRIEF_BELOW else DEPTHS[2] if score >= IN_DEPTH_AT else DEPTHS[1]
    if part["must_cover"] and depth == "brief":
        depth = "normal"
    if part["main"]:
        depth = "in depth"
    return depth


def plan_depth(lecture_id, topic, persona, slides, material):
    """The depth plan: {"parts": [{title, slides, things, yes, score, depth, minutes, words}], "speaking_min",
    "too_long"}, or None (no sub-topics: the caller falls back to an even split)."""
    parts = split_material(topic, slides, material, persona)
    if not parts:
        return None
    things = _slide_things(lecture_id) if slides else {}
    names = [p["title"] for p in parts]

    def state(p):
        if p["slides"]:
            text = "\n\n".join(f"Slide {n}: {slides.get(n) or slides.get(str(n), '')}"
                               + (f"\n(On this slide: {', '.join(things.get(n, []))})" if things.get(n) else "")
                               for n in p["slides"])
        else:
            text = p.get("excerpt", "")
        return {"lesson_topic": topic, "sub_topic": p["title"], "material": text[:MATERIAL_CHARS],
                "other_sub_topics": "; ".join(n for n in names if n != p["title"])}

    with ThreadPoolExecutor(max_workers=6) as pool:
        answers = list(pool.map(_ask_jev, [state(p) for p in parts]))
    for p, yes in zip(parts, answers):
        p["things"] = sorted({t for n in p["slides"] for t in things.get(n, [])})
        p["yes"] = [k for k, v in (yes or {}).items() if v >= YES_AT]
        p["score"] = round(sum((yes or {}).values()) + max(0, len(p["slides"]) - 1), 2)
        p["depth"] = _depth(p["score"], p) if yes is not None else "normal"  # no answer: the middle, not a guess
        p["minutes"] = DEPTH_MIN[p["depth"]]
    speaking = sum(p["minutes"] for p in parts) + sum(FRAMING_MIN.values())
    too_long = speaking > MAX_SPEAKING_MIN
    if too_long:  # shorten every part by the same share, so the plan still fits one sitting
        scale = (MAX_SPEAKING_MIN - sum(FRAMING_MIN.values())) / (speaking - sum(FRAMING_MIN.values()))
        for p in parts:
            p["minutes"] = round(p["minutes"] * scale, 2)
    for p in parts:
        p["words"] = max(60, round(p["minutes"] * config.SPEAKING_WPM))
        p.pop("excerpt", None)
    speaking = round(sum(p["minutes"] for p in parts) + sum(FRAMING_MIN.values()), 1)
    return {"parts": parts, "speaking_min": speaking, "too_long": too_long,
            "framing_words": {k: round(v * config.SPEAKING_WPM) for k, v in FRAMING_MIN.items()}}


def plan_text(plan):
    """The length plan for the lesson writer: one concept section per sub-topic, in order, with its words."""
    f = plan["framing_words"]
    lines = [f'"# Opening" about {f["Opening"]} words.']
    for p in plan["parts"]:
        how = {"brief": "brief: state it clearly with at most one quick illustration",
               "normal": "normal: explain it with one example",
               "in depth": "in depth: build it up step by step, work an example fully, check understanding"}[p["depth"]]
        where = f' (slides {", ".join(map(str, p["slides"]))})' if p["slides"] else ""
        lines.append(f'"# {p["title"]}"{where}: about {p["words"]} words, {how}.')
    lines.append(f'"# Recap" about {f["Recap"]} words; "# Closing" about {f["Closing"]} words.')
    total = sum(p["words"] for p in plan["parts"]) + sum(f.values())
    return ("Write exactly these sections, in this order, with these headings. How long each one is comes from how "
            "the teacher's material treats it, so keep to it: do not stretch a brief part or rush an in-depth one.\n"
            + "\n".join(f"- {l}" for l in lines) + f"\nTotal about {total} words."), total


def section_problems(sections, plan):
    """Concept sections far from their planned length: too long wastes the class's time (and tokens), too short
    rushes it. Matched by heading; a section the writer renamed is matched by its place."""
    planned = {p["title"].lower(): p for p in plan["parts"]}
    concept = [s for s in sections if not FRAMING_RE.match(s["title"])]
    problems = []
    for i, s in enumerate(concept):
        p = planned.get(s["title"].lower()) or (plan["parts"][i] if i < len(plan["parts"]) else None)
        if not p:
            continue
        words = sum(len(it["text"].split()) for it in s["items"] if it["kind"] == "say")
        if words > p["words"] * 1.6:
            problems.append(f'section "{s["title"]}" is {p["depth"]} in the plan (about {p["words"]} spoken words) but has '
                            f"{words}: shorten it to about {p['words']}")
        elif words < p["words"] * MIN_SHARE[p["depth"]]:
            problems.append(f'section "{s["title"]}" is {p["depth"]} in the plan (about {p["words"]} spoken words) but has '
                            f"only {words}: teach it fully, about {p['words']} words")
    return problems


def summary(plan):
    """What Studio shows: each sub-topic's depth and why, and the planned speaking time."""
    return {"speaking_min": plan["speaking_min"], "too_long": plan["too_long"],
            "parts": [{"title": p["title"], "depth": p["depth"], "slides": p["slides"],
                       "why": [LABELS[k] for k in p["yes"]] + ([f"{len(p['slides'])} slides"] if len(p["slides"]) > 1 else [])
                              + (["the lesson's main idea"] if p["main"] else [])}
                      for p in plan["parts"]]}
